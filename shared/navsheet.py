"""ヒアリングシート（xlsx）の読み書き。

先頭シートの契約:

- ``B1``: 種別マーカー ``hearing-sheet``（systemplan / global と同じ。旧 ``navigation-sheet`` も読む）
- 2 行目: 見出し ``項目`` / ``回答``
- 3 行目以降: 設問（A）と回答（B）。回答の改行はセル内改行
- A 列 ``生成指示`` の B セル: Markdown エディタへ渡す処理指示

旧形式（A 列 ``ナビゲーション`` の B セルに Markdown 表）も読む。
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from typing import Any

NAV_LABEL = "ナビゲーション"
INSTRUCTION_LABEL = "生成指示"
ITEM_HEADER = "項目"
VALUE_HEADER = "回答"
MARKER = "hearing-sheet"
LEGACY_MARKERS = ("navigation-sheet",)
SHEET_NAME = "ヒアリングシート"
KIND_LABEL = "種別"


_ROW_RE = re.compile(
    r"^\|\s*(?P<label>(?:\\\||[^|])*?)\s*\|\s*(?P<value>(?:\\\||[^|])*?)\s*\|\s*$"
)


@dataclass
class NavRow:
    label: str
    value: str


@dataclass
class NavSheet:
    rows: list[NavRow] = field(default_factory=list)
    instruction: str = ""


def escape_cell(text: str) -> str:
    return (text or "").replace("|", "\\|").replace("\n", "<br>")


def unescape_cell(text: str) -> str:
    return (text or "").replace("<br>", "\n").replace("\\|", "|")


def table_to_markdown(rows: list[NavRow] | list[dict[str, str]]) -> str:
    lines = ["| 項目 | 値 |", "| --- | --- |"]
    for row in rows:
        if isinstance(row, dict):
            label = str(row.get("label") or "")
            value = str(row.get("value") or "")
        else:
            label = row.label
            value = row.value
        lines.append(f"| {escape_cell(label)} | {escape_cell(value)} |")
    return "\n".join(lines)


def markdown_to_rows(text: str) -> list[NavRow]:
    rows: list[NavRow] = []
    for raw in (text or "").splitlines():
        line = raw.strip()
        if not line.startswith("|"):
            continue
        m = _ROW_RE.match(line)
        if not m:
            continue
        label = unescape_cell(m.group("label")).strip()
        value = unescape_cell(m.group("value")).strip()
        if label in (ITEM_HEADER, "---") or set(label) <= {"-", ":"}:
            continue
        if not label and not value:
            continue
        rows.append(NavRow(label=label, value=value))
    return rows


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).replace("\r\n", "\n").replace("\r", "\n")


def _answer_text(value: str) -> str:
    return _cell_text(value).replace("<br>", "\n")


def parse_workbook(raw: bytes) -> NavSheet:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        rows: list[NavRow] = []
        instruction = ""
        for row in ws.iter_rows(min_col=1, max_col=2, values_only=True):
            label = _cell_text(row[0]).strip() if row else ""
            value = _cell_text(row[1]) if row and len(row) > 1 else ""
            if not label:
                continue
            if label == KIND_LABEL:
                continue
            if label == NAV_LABEL:
                rows.extend(markdown_to_rows(value))
                continue
            if label == INSTRUCTION_LABEL:
                instruction = _answer_text(value).strip()
                continue
            if label == ITEM_HEADER and value.strip() in (VALUE_HEADER, "値", ""):
                continue
            rows.append(NavRow(label=label, value=_answer_text(value).strip()))
        return NavSheet(rows=rows, instruction=instruction)
    finally:
        wb.close()


def write_workbook(sheet: NavSheet) -> bytes:
    import openpyxl
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = SHEET_NAME
    header = Font(bold=True)
    wrap = Alignment(wrap_text=True, vertical="top")

    ws["A1"] = KIND_LABEL
    ws["B1"] = MARKER
    ws["A2"] = ITEM_HEADER
    ws["B2"] = VALUE_HEADER
    ws["A1"].font = header
    ws["A2"].font = header
    ws["B2"].font = header

    start = 3
    for i, row in enumerate(sheet.rows):
        r = start + i
        ws.cell(r, 1, row.label or "")
        answer = _answer_text(row.value)
        cell = ws.cell(r, 2, answer)
        cell.alignment = wrap
        lines = max(1, answer.count("\n") + 1)
        ws.row_dimensions[r].height = min(220, max(18, 15 * lines))

    inst_row = start + len(sheet.rows)
    ws.cell(inst_row, 1, INSTRUCTION_LABEL).font = header
    inst = ws.cell(inst_row, 2, _answer_text(sheet.instruction or ""))
    inst.alignment = wrap
    ws.row_dimensions[inst_row].height = 80

    ws.column_dimensions[get_column_letter(1)].width = 24
    ws.column_dimensions[get_column_letter(2)].width = 80
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def empty_workbook() -> bytes:
    return write_workbook(NavSheet(rows=[], instruction=""))


def slug_for(label: str, index: int) -> str:
    text = re.sub(r"[^\w]+", "-", (label or "").strip(), flags=re.UNICODE)
    text = text.strip("-").lower() or f"item-{index}"
    return text[:40]
