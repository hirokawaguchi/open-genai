"""Markdown 章を docx / html / pptx / txt / md へ変換する（/compose の新形式）。"""

from __future__ import annotations

import base64
import html
import io
import re
from datetime import datetime
from typing import Any

from app.dads import (
    ACCENT,
    BODY,
    DOCX_FONT_MONO,
    FONT,
    FONT_MONO,
    INK,
    MUTED,
    ON_ACCENT,
    PAPER,
    RULE,
    SURFACE,
    apply_docx_theme,
    bottom_border,
    hex_of,
    left_border,
    set_cell_borders,
    set_table_full_width,
    shade_cell,
    shade_paragraph,
    shade_run,
    style_run,
)

# 画像のみの行（ブロック画像）。
_IMAGE_LINE_RE = re.compile(r"^!\[[^\]]*\]\(\s*<?([^)>\s]+)>?(?:\s+\"[^\"]*\")?\s*\)$")
_INLINE_IMAGE_RE = re.compile(r"!\[([^\]]*)\]\(\s*<?([^)>\s]+)>?(?:\s+\"[^\"]*\")?\s*\)")
_LINK_RE = re.compile(r"\[([^\]]+)\]\([^)]+\)")
_BOLD_RE = re.compile(r"(\*\*|__)(.+?)\1")
_ITALIC_RE = re.compile(r"(?<!\*)\*(?!\*)(.+?)(?<!\*)\*(?!\*)")
_HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")
_BULLET_RE = re.compile(r"^(\s*)[-*]\s+(.*)$")
_NUMBER_RE = re.compile(r"^(\s*)(\d+)[.)]\s+(.*)$")
_QUOTE_RE = re.compile(r"^>\s?(.*)$")
_HR_RE = re.compile(r"^(?:-{3,}|\*{3,}|_{3,})\s*$")
_TASK_RE = re.compile(r"^\[([ xX])\]\s+(.*)$")
_TABLE_LINE_RE = re.compile(r"^\s*\|.+\|\s*$")
_TABLE_SEP_CELL_RE = re.compile(r"^:?-+:?$")
_INLINE_RE = re.compile(
    r"!\[(?P<img_alt>[^\]]*)\]\(\s*<?(?P<img_src>[^)>\s]+)>?(?:\s+\"[^\"]*\")?\s*\)"
    r"|`(?P<code>[^`]+)`"
    r"|\*\*(?P<bold>.+?)\*\*"
    r"|__(?P<bold2>.+?)__"
    r"|~~(?P<strike>.+?)~~"
    r"|\[(?P<link_text>[^\]]+)\]\((?P<link_href>[^)]+)\)"
    r"|(?<!\*)\*(?P<italic>(?:(?!\*).)+?)\*(?!\*)"
)

SUPPORTED_FORMATS = ("docx", "html", "pptx", "txt", "md")

def _css_hex(rgb: tuple[int, int, int]) -> str:
    return f"#{hex_of(rgb)}"


# 庁内情報伝達向けの縦スクロール文書。色は DADS トークン、フォントはブラウザ表示に
# 適した system スタック＋和文フォールバック（Noto Sans JP は Web と揃う）。
_HTML_CSS = f"""
:root {{
  --ink: {_css_hex(INK)};
  --body: {_css_hex(BODY)};
  --muted: {_css_hex(MUTED)};
  --rule: {_css_hex(RULE)};
  --paper: {_css_hex(PAPER)};
  --surface: {_css_hex(SURFACE)};
  --accent: {_css_hex(ACCENT)};
}}
* {{ box-sizing: border-box; }}
@page {{ margin: 18mm 16mm; }}
html {{ background: var(--surface); }}
body {{
  margin: 0;
  color: var(--body);
  font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", "{FONT}", "Hiragino Sans", "Yu Gothic", sans-serif;
  font-size: 16px;
  line-height: 1.75;
}}
main.doc {{
  max-width: 46rem;
  margin: 2.5rem auto 4rem;
  padding: 2.4rem 2.6rem 3rem;
  background: var(--paper);
  border: 1px solid var(--rule);
  border-radius: 12px;
  box-shadow: 0 10px 28px rgba(0, 0, 0, 0.06);
}}
.doc-header {{ margin: 0 0 1.6rem; }}
.doc-header h1 {{
  margin: 0;
  padding-bottom: 0.6rem;
  border-bottom: 3px solid var(--accent);
  font-size: 1.7rem;
  line-height: 1.4;
  color: var(--ink);
}}
.doc-meta {{ margin: 0.5rem 0 0; color: var(--muted); font-size: 0.85rem; }}
h1, h2, h3, h4 {{ color: var(--ink); font-weight: 700; line-height: 1.5; }}
h2 {{
  font-size: 1.3rem;
  margin: 2rem 0 0.7rem;
  padding-bottom: 0.35rem;
  border-bottom: 1px solid var(--rule);
}}
h3 {{ font-size: 1.12rem; margin: 1.5rem 0 0.5rem; }}
h4 {{ font-size: 1rem; margin: 1.2rem 0 0.4rem; }}
p {{ margin: 0 0 0.9rem; }}
ul, ol {{ margin: 0 0 1rem 1.4rem; padding: 0; }}
li {{ margin: 0.25rem 0; }}
li.task {{ list-style: none; margin-left: -1.1rem; }}
a {{ color: var(--accent); }}
strong {{ font-weight: 700; }}
s {{ color: var(--muted); }}
figure {{ margin: 1.2rem 0; }}
img {{ max-width: 100%; height: auto; display: block; margin: 0 auto; }}
hr {{ border: none; border-top: 1px solid var(--rule); margin: 1.8rem 0; }}
code, pre {{
  font-family: "{FONT_MONO}", ui-monospace, "SFMono-Regular", Menlo, monospace;
  font-size: 0.88em;
}}
code {{ background: var(--surface); padding: 0.1em 0.35em; border-radius: 4px; }}
pre {{
  background: var(--surface);
  padding: 0.85rem 1rem;
  overflow: auto;
  line-height: 1.55;
  border: 1px solid var(--rule);
  border-radius: 6px;
}}
pre code {{ background: none; padding: 0; }}
blockquote {{
  margin: 0 0 1rem;
  padding: 0.3rem 0 0.3rem 1rem;
  border-left: 4px solid var(--accent);
  color: var(--muted);
}}
.table-wrap {{ overflow-x: auto; margin: 1rem 0 1.2rem; }}
table {{ border-collapse: collapse; width: 100%; font-size: 0.95em; }}
th, td {{ border: 1px solid var(--rule); padding: 0.45rem 0.7rem; text-align: left; vertical-align: top; }}
th {{ background: var(--surface); font-weight: 700; }}
nav.toc {{
  margin: 0 0 1.8rem;
  padding: 1rem 1.2rem;
  background: var(--surface);
  border: 1px solid var(--rule);
  border-radius: 8px;
}}
nav.toc .toc-title {{ margin: 0 0 0.5rem; font-weight: 700; color: var(--ink); }}
nav.toc ul {{ margin: 0; padding: 0; list-style: none; }}
nav.toc li {{ margin: 0.2rem 0; }}
nav.toc li.toc-l3 {{ margin-left: 1.2rem; font-size: 0.94em; }}
nav.toc a {{ color: var(--accent); text-decoration: none; }}
nav.toc a:hover {{ text-decoration: underline; }}
.figure-missing {{ color: var(--muted); font-style: italic; }}
footer.meta {{ margin-top: 2.5rem; padding-top: 1rem; border-top: 1px solid var(--rule); color: var(--muted); font-size: 0.8rem; }}
@media (max-width: 720px) {{
  main.doc {{ margin: 0; border: none; border-radius: 0; padding: 1.4rem 1.1rem 2rem; box-shadow: none; }}
  body {{ font-size: 15px; }}
  .doc-header h1 {{ font-size: 1.45rem; }}
}}
@media print {{
  html {{ background: #fff; }}
  main.doc {{ margin: 0; max-width: none; border: none; border-radius: 0; box-shadow: none; padding: 0; }}
  nav.toc {{ background: #fff; }}
  h1, h2, h3, h4 {{ break-after: avoid; }}
  table, figure, pre, blockquote {{ break-inside: avoid; }}
}}
""".strip()


def normalize_format(raw: Any) -> str:
    fmt = str(raw or "docx").strip().lower().lstrip(".")
    return fmt if fmt in SUPPORTED_FORMATS else "docx"


def _rel_of(path: str) -> str:
    return path.replace("\\", "/").lstrip("/")


def _join_sections(sections: list[dict[str, Any]]) -> str:
    parts: list[str] = []
    for sec in sections:
        content = str(sec.get("content") or "").strip()
        if content:
            parts.append(content)
    return "\n\n".join(parts)


def _data_uri(rel: str, data: bytes) -> str:
    ext = rel.rsplit(".", 1)[-1].lower() if "." in rel else ""
    mime = {
        "png": "image/png",
        "jpg": "image/jpeg",
        "jpeg": "image/jpeg",
        "gif": "image/gif",
        "webp": "image/webp",
        "svg": "image/svg+xml",
    }.get(ext, "application/octet-stream")
    return f"data:{mime};base64,{base64.b64encode(data).decode('ascii')}"


def markdown_to_md(sections: list[dict[str, Any]]) -> bytes:
    """章本文を空行区切りで連結する（ソースそのもの）。"""
    body = _join_sections(sections)
    if body and not body.endswith("\n"):
        body += "\n"
    return body.encode("utf-8")


def markdown_to_txt(sections: list[dict[str, Any]]) -> bytes:
    """Markdown 記法を除いたプレーンテキスト。画像・Mermaid はプレースホルダ。"""
    lines_out: list[str] = []
    in_code = False
    code_lang = ""
    for raw in _join_sections(sections).splitlines():
        stripped = raw.strip()
        if stripped.startswith("```"):
            if in_code:
                if code_lang == "mermaid":
                    lines_out.append("[図]")
                in_code = False
                code_lang = ""
            else:
                in_code = True
                code_lang = stripped[3:].strip().lower()
            continue
        if in_code:
            if code_lang != "mermaid":
                lines_out.append(raw)
            continue
        m = _IMAGE_LINE_RE.match(stripped)
        if m:
            lines_out.append(f"[画像: {_rel_of(m.group(1))}]")
            continue
        line = _INLINE_IMAGE_RE.sub(lambda mm: f"[画像: {_rel_of(mm.group(2))}]", raw)
        line = _LINK_RE.sub(r"\1", line)
        line = _BOLD_RE.sub(r"\2", line)
        line = _ITALIC_RE.sub(r"\1", line)
        hm = _HEADING_RE.match(line)
        if hm:
            line = hm.group(2).strip()
        bm = _BULLET_RE.match(line)
        if bm:
            line = f"{bm.group(1)}・ {bm.group(2)}"
        lines_out.append(line.rstrip())
    text = "\n".join(lines_out).strip() + "\n"
    return text.encode("utf-8")


def _safe_href(href: str) -> str | None:
    """リンク先を検証する。javascript: 等の危険なスキームは弾く。

    許可: http / https / mailto / tel / 相対・アンカー（スキーム無し）。
    """
    h = (href or "").strip()
    if not h:
        return None
    scheme = ""
    if ":" in h.split("/", 1)[0] and not h.startswith(("/", "#", "?", ".")):
        scheme = h.split(":", 1)[0].strip().lower()
    if scheme and scheme not in ("http", "https", "mailto", "tel"):
        return None
    return h


def _inline_html(text: str, assets: dict[str, bytes]) -> str:
    """プレビュー相当のインライン装飾（太字・斜体・取消線・コード・リンク・画像）。

    docx の `_add_inline_runs` と同じ `_INLINE_RE` を使い、機能を揃える。
    """
    out: list[str] = []
    last = 0
    for m in _INLINE_RE.finditer(text):
        if m.start() > last:
            out.append(html.escape(text[last : m.start()]))
        g = m.groupdict()
        if g.get("img_src"):
            rel = _rel_of(g["img_src"])
            data = assets.get(rel)
            alt = html.escape(g.get("img_alt") or "")
            if data:
                out.append(f'<img src="{_data_uri(rel, data)}" alt="{alt}">')
            else:
                out.append(f'<span class="figure-missing">[画像: {html.escape(rel)}]</span>')
        elif g.get("code"):
            out.append(f"<code>{html.escape(g['code'])}</code>")
        elif g.get("bold") or g.get("bold2"):
            out.append(f"<strong>{html.escape(g.get('bold') or g.get('bold2') or '')}</strong>")
        elif g.get("strike"):
            out.append(f"<s>{html.escape(g['strike'])}</s>")
        elif g.get("link_text"):
            label = html.escape(g["link_text"])
            safe = _safe_href(g.get("link_href") or "")
            if safe:
                out.append(
                    f'<a href="{html.escape(safe, quote=True)}"'
                    ' rel="noopener noreferrer">'
                    f"{label}</a>"
                )
            else:
                out.append(label)
        elif g.get("italic"):
            out.append(f"<em>{html.escape(g['italic'])}</em>")
        last = m.end()
    if last < len(text):
        out.append(html.escape(text[last:]))
    return "".join(out)


def _table_cells(line: str) -> list[str]:
    s = line.strip()
    if s.startswith("|"):
        s = s[1:]
    if s.endswith("|"):
        s = s[:-1]
    return [c.strip() for c in s.split("|")]


def _is_table_sep(cells: list[str]) -> bool:
    return bool(cells) and all(
        _TABLE_SEP_CELL_RE.match(c.replace(" ", "")) for c in cells
    )


def parse_gfm_table(lines: list[str]) -> dict[str, list[list[str]] | list[str]] | None:
    """GFM 表を headers / rows に分解する。区切り行は捨てる。"""
    raw = [ln for ln in lines if ln.strip()]
    if len(raw) < 2:
        return None
    parsed = [_table_cells(ln) for ln in raw]
    header = parsed[0]
    if not header or not any(header):
        return None
    width = len(header)
    rows: list[list[str]] = []
    for cells in parsed[1:]:
        if _is_table_sep(cells):
            continue
        padded = cells + [""] * width
        rows.append(padded[:width])
    if not rows:
        return None
    return {"headers": header, "rows": rows}


def _table_html(block: list[str], assets: dict[str, bytes]) -> str:
    rows = [_table_cells(ln) for ln in block]
    header = rows[0]
    body = [r for r in rows[1:] if not _is_table_sep(r)]
    width = len(header)
    parts = ["<table>", "<thead>", "<tr>"]
    for cell in header:
        parts.append(f"<th>{_inline_html(cell, assets)}</th>")
    parts.extend(["</tr>", "</thead>", "<tbody>"])
    for row in body:
        padded = row + [""] * width
        parts.append("<tr>")
        for cell in padded[:width]:
            parts.append(f"<td>{_inline_html(cell, assets)}</td>")
        parts.append("</tr>")
    parts.extend(["</tbody>", "</table>"])
    return "".join(parts)


def _toc_slug(seq: int) -> str:
    """目次アンカー用の安定した id（非 ASCII を避け URL エンコード不要にする）。"""
    return f"sec-{seq}"


def markdown_to_html(
    name: str, sections: list[dict[str, Any]], assets: dict[str, bytes] | None = None
) -> bytes:
    """庁内情報伝達向けの自己完結・縦スクロール HTML（CSS 埋め込み・画像 data URI）。

    以後編集しない最終成果物として、情報が伝わることを優先する。見出しから目次を作り、
    Markdown の装飾・表・リスト・引用・コード・水平線・画像を docx 相当でカバーする。
    """
    assets = assets or {}
    chunks: list[str] = []
    headings: list[tuple[int, str, str]] = []  # (level, id, text)
    heading_seq = 0
    in_code = False
    code_lang = ""
    code_lines: list[str] = []
    list_kind: str | None = None  # "ul" | "ol"

    def close_list() -> None:
        nonlocal list_kind
        if list_kind:
            chunks.append(f"</{list_kind}>")
            list_kind = None

    def open_list(kind: str) -> None:
        nonlocal list_kind
        if list_kind != kind:
            close_list()
            chunks.append(f"<{kind}>")
            list_kind = kind

    def flush_code() -> None:
        nonlocal in_code, code_lang, code_lines
        if code_lang == "mermaid":
            chunks.append(
                '<p class="figure-missing">【Mermaid 図（画像未変換のためソースを表示）】</p>'
            )
        body = html.escape("\n".join(code_lines))
        chunks.append(f"<pre><code>{body}</code></pre>")
        in_code, code_lang, code_lines = False, "", []

    lines = _join_sections(sections).splitlines()
    i = 0
    while i < len(lines):
        raw_line = lines[i]
        stripped = raw_line.strip()
        if stripped.startswith("```"):
            if in_code:
                flush_code()
            else:
                close_list()
                in_code, code_lang, code_lines = True, stripped[3:].strip().lower(), []
            i += 1
            continue
        if in_code:
            code_lines.append(raw_line)
            i += 1
            continue
        if _TABLE_LINE_RE.match(stripped):
            block = [raw_line]
            j = i + 1
            while j < len(lines) and _TABLE_LINE_RE.match(lines[j].strip()):
                block.append(lines[j])
                j += 1
            if len(block) >= 2 and _is_table_sep(_table_cells(block[1])):
                close_list()
                chunks.append(f'<div class="table-wrap">{_table_html(block, assets)}</div>')
                i = j
                continue
        if not stripped:
            close_list()
            i += 1
            continue
        hm = _HEADING_RE.match(stripped)
        if hm:
            close_list()
            level = min(len(hm.group(1)) + 1, 4)
            text = hm.group(2).strip()
            heading_seq += 1
            hid = _toc_slug(heading_seq)
            if level in (2, 3):
                headings.append((level, hid, text))
            chunks.append(f'<h{level} id="{hid}">{_inline_html(text, assets)}</h{level}>')
            i += 1
            continue
        if _HR_RE.match(stripped):
            close_list()
            chunks.append("<hr>")
            i += 1
            continue
        m = _IMAGE_LINE_RE.match(stripped)
        if m:
            close_list()
            rel = _rel_of(m.group(1))
            data = assets.get(rel)
            if data:
                chunks.append(
                    f'<figure><img src="{_data_uri(rel, data)}" alt=""></figure>'
                )
            else:
                chunks.append(f'<p class="figure-missing">[画像: {html.escape(rel)}]</p>')
            i += 1
            continue
        bm = _BULLET_RE.match(raw_line)
        if bm:
            open_list("ul")
            text = bm.group(2)
            task = _TASK_RE.match(text)
            if task:
                mark = "☑" if task.group(1).lower() == "x" else "☐"
                chunks.append(
                    f'<li class="task">{mark} {_inline_html(task.group(2), assets)}</li>'
                )
            else:
                chunks.append(f"<li>{_inline_html(text, assets)}</li>")
            i += 1
            continue
        nm = _NUMBER_RE.match(raw_line)
        if nm:
            open_list("ol")
            chunks.append(f"<li>{_inline_html(nm.group(3), assets)}</li>")
            i += 1
            continue
        qm = _QUOTE_RE.match(stripped)
        if qm:
            close_list()
            quote_lines = [qm.group(1)]
            j = i + 1
            while j < len(lines):
                q2 = _QUOTE_RE.match(lines[j].strip())
                if not q2 or not lines[j].strip():
                    break
                quote_lines.append(q2.group(1))
                j += 1
            body = "<br>".join(_inline_html(q, assets) for q in quote_lines)
            chunks.append(f"<blockquote><p>{body}</p></blockquote>")
            i = j
            continue
        close_list()
        chunks.append(f"<p>{_inline_html(stripped, assets)}</p>")
        i += 1
    if in_code:
        flush_code()
    close_list()

    # 目次: 見出し（h2/h3）が十分あるときだけ出す。
    toc_html = ""
    if len(headings) >= 3:
        items = "".join(
            f'<li class="toc-l{level}"><a href="#{hid}">{html.escape(text)}</a></li>'
            for level, hid, text in headings
        )
        toc_html = (
            '<nav class="toc" aria-label="目次">'
            "<p class=\"toc-title\">目次</p>"
            f"<ul>{items}</ul></nav>\n"
        )

    generated = datetime.now().strftime("%Y-%m-%d")
    header_html = (
        '<header class="doc-header">'
        f"<h1>{html.escape(name)}</h1>"
        f'<p class="doc-meta">作成日: {generated}</p>'
        "</header>\n"
    )
    footer_html = (
        f'\n<footer class="meta">作成日: {generated}</footer>'
    )

    doc = (
        "<!DOCTYPE html>\n"
        '<html lang="ja">\n<head>\n<meta charset="utf-8">\n'
        '<meta name="viewport" content="width=device-width, initial-scale=1">\n'
        f"<title>{html.escape(name)}</title>\n"
        f"<style>\n{_HTML_CSS}\n</style>\n</head>\n<body>\n"
        '<main class="doc">\n'
        + header_html
        + toc_html
        + "\n".join(chunks)
        + footer_html
        + "\n</main>\n</body>\n</html>\n"
    )
    return doc.encode("utf-8")


# PPTX は dads のトークンをそのまま使う（docx / html と揃える）。
_PPTX_INK = INK
_PPTX_BODY = BODY
_PPTX_MUTED = MUTED
_PPTX_RULE = RULE
_PPTX_PAPER = PAPER
_PPTX_SURFACE = SURFACE
_PPTX_ACCENT = ACCENT
_PPTX_ON_ACCENT = ON_ACCENT
_PPTX_FONT = FONT


def _rgb(rgb: tuple[int, int, int]) -> Any:
    from pptx.dml.color import RGBColor

    return RGBColor(*rgb)


def _no_line(shape: Any) -> None:
    shape.line.fill.background()


def _solid(shape: Any, rgb: tuple[int, int, int]) -> None:
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(rgb)
    _no_line(shape)


def _set_run_font(
    run: Any,
    *,
    size: Any,
    color: tuple[int, int, int],
    bold: bool = False,
    name: str = _PPTX_FONT,
) -> None:
    from lxml import etree
    from pptx.oxml.ns import qn

    run.font.size = size
    run.font.bold = bold
    run.font.color.rgb = _rgb(color)
    run.font.name = name
    r_pr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = r_pr.find(qn(tag))
        if el is None:
            el = etree.SubElement(r_pr, qn(tag))
        el.set("typeface", name)


def _fill_slide(slide: Any, rgb: tuple[int, int, int]) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = _rgb(rgb)


def _add_rect(slide: Any, left: Any, top: Any, width: Any, height: Any, rgb: tuple[int, int, int]) -> Any:
    from pptx.enum.shapes import MSO_SHAPE

    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    _solid(shape, rgb)
    return shape


def _style_paragraph(
    paragraph: Any,
    text: str,
    *,
    size: Any,
    color: tuple[int, int, int],
    bold: bool = False,
    space_after: Any | None = None,
    space_before: Any | None = None,
) -> None:
    from pptx.enum.text import PP_ALIGN

    paragraph.text = text
    paragraph.alignment = PP_ALIGN.LEFT
    if space_after is not None:
        paragraph.space_after = space_after
    if space_before is not None:
        paragraph.space_before = space_before
    for run in paragraph.runs:
        _set_run_font(run, size=size, color=color, bold=bold)


def _add_title_slide(prs: Any, title: str) -> None:
    """表紙。DADS の Blue 900 ヘッダー＋白地。装飾色は使わない。"""
    from pptx.util import Emu, Inches, Pt

    blank = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank)
    _fill_slide(slide, _PPTX_PAPER)
    _add_rect(slide, Inches(0), Inches(0), prs.slide_width, Inches(1.05), _PPTX_ACCENT)
    label = slide.shapes.add_textbox(Inches(0.7), Inches(0.32), Inches(12.0), Inches(0.45))
    _style_paragraph(label.text_frame.paragraphs[0], "資料", size=Pt(14), color=_PPTX_ON_ACCENT, bold=True)

    box = slide.shapes.add_textbox(Inches(0.7), Inches(2.35), Inches(12.0), Inches(2.4))
    box.text_frame.word_wrap = True
    _style_paragraph(box.text_frame.paragraphs[0], title, size=Pt(36), color=_PPTX_INK, bold=True)
    _add_rect(slide, Inches(0.7), Inches(5.0), Inches(3.2), Inches(0.04), _PPTX_ACCENT)
    _add_rect(
        slide, Inches(0), Emu(prs.slide_height - Inches(0.28)), prs.slide_width, Inches(0.28), _PPTX_SURFACE
    )


def _add_content_chrome(slide: Any, prs: Any, title: str) -> None:
    from pptx.util import Inches, Pt

    _fill_slide(slide, _PPTX_PAPER)
    _add_rect(slide, Inches(0), Inches(0), prs.slide_width, Inches(0.08), _PPTX_ACCENT)
    box = slide.shapes.add_textbox(Inches(0.7), Inches(0.28), Inches(12.0), Inches(0.85))
    box.text_frame.word_wrap = True
    _style_paragraph(box.text_frame.paragraphs[0], title, size=Pt(24), color=_PPTX_INK, bold=True)
    _add_rect(slide, Inches(0.7), Inches(1.16), Inches(11.9), Inches(0.012), _PPTX_RULE)


def _add_slide_footers(prs: Any) -> None:
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Emu, Inches, Pt

    total = len(prs.slides)
    for i, slide in enumerate(prs.slides, start=1):
        if i == 1:
            continue
        _add_rect(
            slide,
            Inches(0.7),
            Emu(prs.slide_height - Inches(0.42)),
            Inches(11.9),
            Inches(0.01),
            _PPTX_RULE,
        )
        box = slide.shapes.add_textbox(
            Inches(11.4), Emu(prs.slide_height - Inches(0.38)), Inches(1.3), Inches(0.3)
        )
        p = box.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.RIGHT
        run = p.add_run()
        run.text = f"{i} / {total}"
        _set_run_font(run, size=Pt(10), color=_PPTX_MUTED)


def markdown_to_pptx(
    name: str, sections: list[dict[str, Any]], assets: dict[str, bytes] | None = None
) -> bytes:
    """`#` / `##` 見出しでスライドを分け、本文と画像を配置する。"""
    from pptx import Presentation
    from pptx.util import Inches, Pt

    assets = assets or {}
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    _add_title_slide(prs, name)

    slide = None
    body_lines: list[tuple[str, bool]] = []
    table_buf: list[str] = []
    text_flushed = False
    in_code = False
    code_lang = ""
    # notes-first と揃え、フォールバックでもスライドの原文をノートに残す。
    note_heading = name
    note_buf: list[str] = []

    def flush_notes() -> None:
        if slide is None:
            return
        original = "\n".join(x for x in note_buf if x.strip()).strip()
        if not original:
            return
        from app.pptx_layouts import _add_notes

        _add_notes(slide, f"元の本文（{note_heading}）\n{original}")

    def flush_table() -> None:
        nonlocal table_buf, text_flushed
        parsed = parse_gfm_table(table_buf)
        table_buf = []
        if not parsed or slide is None:
            return
        from pptx.util import Inches

        from app.pptx_layouts import _add_table

        rows = [list(parsed["headers"]), *[list(r) for r in parsed["rows"]]]
        top = Inches(4.55) if text_flushed else Inches(1.4)
        height = Inches(2.15) if text_flushed else Inches(5.1)
        _add_table(slide, Inches(0.7), top, Inches(11.9), height, rows)
        text_flushed = True

    def flush_text(*, narrow: bool = False) -> None:
        nonlocal body_lines, text_flushed
        if slide is None or not body_lines:
            body_lines = []
            return
        # 同じ座標に重ねない（本文のあと画像、そのあとまた本文、で文字が重なっていた）。
        top = Inches(4.55) if text_flushed else Inches(1.4)
        height = Inches(2.15) if text_flushed else Inches(5.3)
        width = Inches(6.6) if narrow else Inches(11.9)
        box = slide.shapes.add_textbox(Inches(0.7), top, width, height)
        tf = box.text_frame
        tf.word_wrap = True
        for i, (line, is_bullet) in enumerate(body_lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            text = f"•   {line}" if is_bullet else line
            _style_paragraph(p, text, size=Pt(16), color=_PPTX_BODY, space_after=Pt(8), space_before=Pt(0))
            p.line_spacing = 1.35
        body_lines = []
        text_flushed = True

    def new_slide(title: str) -> None:
        nonlocal slide, text_flushed, note_heading, note_buf
        flush_text()
        flush_table()
        flush_notes()  # 直前スライドの原文をノートへ
        slide = prs.slides.add_slide(blank)
        _add_content_chrome(slide, prs, title)
        text_flushed = False
        note_heading = title
        note_buf = []

    def add_image(rel: str) -> None:
        data = assets.get(rel)
        if slide is None:
            new_slide(name)
        if not data:
            body_lines.append((f"[画像: {rel}]", False))
            return
        narrow = bool(body_lines)
        flush_text(narrow=narrow)
        stream = io.BytesIO(data)
        try:
            if narrow:
                slide.shapes.add_picture(stream, Inches(7.6), Inches(1.45), width=Inches(5.0))
            else:
                slide.shapes.add_picture(stream, Inches(1.8), Inches(1.5), width=Inches(9.6))
        except Exception:  # noqa: BLE001
            body_lines.append((f"[画像: {rel}]", False))

    for raw_line in _join_sections(sections).splitlines():
        stripped = raw_line.strip()
        if stripped.startswith("```"):
            if in_code:
                in_code = False
                code_lang = ""
            else:
                in_code = True
                code_lang = stripped[3:].strip().lower()
            continue
        if in_code:
            if code_lang != "mermaid":
                body_lines.append((raw_line, False))
                note_buf.append(raw_line)
            continue
        if not stripped:
            flush_table()
            continue
        hm = _HEADING_RE.match(stripped)
        if hm and len(hm.group(1)) <= 2:
            new_slide(hm.group(2).strip())
            continue
        if hm:
            if slide is None:
                new_slide(name)
            flush_table()
            body_lines.append((hm.group(2).strip(), False))
            note_buf.append(hm.group(2).strip())
            continue
        if _TABLE_LINE_RE.match(stripped):
            if slide is None:
                new_slide(name)
            table_buf.append(stripped)
            note_buf.append(stripped)
            continue
        flush_table()
        m = _IMAGE_LINE_RE.match(stripped)
        if m:
            add_image(_rel_of(m.group(1)))
            continue
        if slide is None:
            new_slide(name)
        bm = _BULLET_RE.match(raw_line)
        text = bm.group(2) if bm else stripped
        text = _INLINE_IMAGE_RE.sub("", text).strip()
        text = _BOLD_RE.sub(r"\2", text)
        if text:
            body_lines.append((text, bool(bm)))
            note_buf.append(f"・ {text}" if bm else text)
    flush_text()
    flush_table()
    flush_notes()
    _add_slide_footers(prs)

    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


def _table_alignments(block: list[str]) -> list[str]:
    """区切り行から left / center / right を取る。無ければ left。"""
    if len(block) < 2:
        return []
    seps = _table_cells(block[1])
    if not _is_table_sep(seps):
        return []
    aligns: list[str] = []
    for cell in seps:
        s = cell.replace(" ", "")
        left = s.startswith(":")
        right = s.endswith(":")
        if left and right:
            aligns.append("center")
        elif right:
            aligns.append("right")
        else:
            aligns.append("left")
    return aligns


def _add_hyperlink(paragraph: Any, text: str, url: str, *, size: Any = None) -> None:
    from docx.opc.constants import RELATIONSHIP_TYPE
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    part = paragraph.part
    r_id = part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = paragraph.add_run(text)
    style_run(run, size=size, color=ACCENT)
    run.underline = True
    r_el = run._r
    r_el.getparent().remove(r_el)
    hyperlink.append(r_el)
    paragraph._p.append(hyperlink)


def _add_inline_runs(
    paragraph: Any,
    text: str,
    assets: dict[str, bytes],
    *,
    size: Any = None,
    color: tuple[int, int, int] = BODY,
) -> None:
    """プレビュー相当のインライン装飾（太字・斜体・取消線・コード・リンク・画像）。"""
    from docx.shared import Cm, Pt

    last = 0
    for m in _INLINE_RE.finditer(text):
        if m.start() > last:
            style_run(paragraph.add_run(text[last : m.start()]), size=size, color=color)
        g = m.groupdict()
        if g.get("img_src"):
            rel = _rel_of(g["img_src"])
            data = assets.get(rel)
            if data:
                try:
                    paragraph.add_run().add_picture(io.BytesIO(data), width=Cm(12))
                except Exception:  # noqa: BLE001
                    style_run(paragraph.add_run(f"[画像: {rel}]"), size=Pt(10), color=MUTED, italic=True)
            else:
                style_run(paragraph.add_run(f"[画像: {rel}]"), size=Pt(10), color=MUTED, italic=True)
        elif g.get("code"):
            run = paragraph.add_run(g["code"])
            style_run(run, name=DOCX_FONT_MONO, size=Pt(10) if size is None else size, color=color)
            shade_run(run)
        elif g.get("bold") or g.get("bold2"):
            style_run(
                paragraph.add_run(g.get("bold") or g.get("bold2") or ""),
                size=size,
                color=color,
                bold=True,
            )
        elif g.get("strike"):
            run = paragraph.add_run(g["strike"])
            style_run(run, size=size, color=color)
            run.font.strike = True
        elif g.get("link_text"):
            href = (g.get("link_href") or "").strip()
            if href:
                _add_hyperlink(paragraph, g["link_text"], href, size=size)
            else:
                style_run(paragraph.add_run(g["link_text"]), size=size, color=ACCENT)
        elif g.get("italic"):
            style_run(paragraph.add_run(g["italic"]), size=size, color=color, italic=True)
        last = m.end()
    if last < len(text):
        style_run(paragraph.add_run(text[last:]), size=size, color=color)


def _add_code_block_docx(doc: Any, lang: str, lines: list[str]) -> None:
    from docx.shared import Pt

    if lang == "mermaid":
        note = doc.add_paragraph()
        style_run(
            note.add_run("【Mermaid 図（画像未変換のためソースを表示）】"),
            size=Pt(10),
            color=MUTED,
            italic=True,
        )
    para = doc.add_paragraph()
    shade_paragraph(para)
    para.paragraph_format.line_spacing = 1.45
    para.paragraph_format.space_after = Pt(10)
    style_run(para.add_run("\n".join(lines)), name=DOCX_FONT_MONO, size=Pt(10), color=BODY)


def _add_hr_docx(doc: Any) -> None:
    from docx.shared import Pt

    para = doc.add_paragraph()
    para.paragraph_format.space_before = Pt(6)
    para.paragraph_format.space_after = Pt(10)
    bottom_border(para._p.get_or_add_pPr(), RULE, sz="8", space="1")


def _add_gfm_table_docx(doc: Any, block: list[str], assets: dict[str, bytes]) -> None:
    """GFM 表をネイティブ Word 表にする（罫線・見出し行はプレビューに合わせる）。"""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt

    parsed = parse_gfm_table(block)
    headers = list(parsed["headers"]) if parsed else _table_cells(block[0])
    rows = [list(r) for r in parsed["rows"]] if parsed else []
    if not headers:
        return
    width = len(headers)
    aligns = _table_alignments(block)
    table = doc.add_table(rows=1 + len(rows), cols=width)
    table.autofit = True
    set_table_full_width(table)
    col_w = Cm(16.0 / max(width, 1))
    align_map = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
    }

    def fill(cell: Any, text: str, *, header: bool, col: int) -> None:
        cell.text = ""
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.35
        if col < len(aligns):
            p.alignment = align_map.get(aligns[col], WD_ALIGN_PARAGRAPH.LEFT)
        _add_inline_runs(p, text, assets, size=Pt(11))
        if header:
            for run in p.runs:
                run.bold = True
            shade_cell(cell, SURFACE)
        set_cell_borders(cell)
        cell.width = col_w

    for c, text in enumerate(headers):
        fill(table.cell(0, c), text, header=True, col=c)
    for r, row in enumerate(rows, start=1):
        padded = row + [""] * width
        for c, text in enumerate(padded[:width]):
            fill(table.cell(r, c), text, header=False, col=c)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(8)


def markdown_to_docx(
    name: str, sections: list[dict[str, Any]], assets: dict[str, bytes] | None = None
) -> bytes:
    """章（Markdown）をプレビュー相当の表現で .docx にする（python-docx）。

    GFM 表はネイティブ表。太字・斜体・取消線・インラインコード・リンク・番号付きリスト・
    引用・水平線も反映する。画像は assets から埋め込む。
    """
    from docx import Document
    from docx.shared import Cm, Pt

    assets = assets or {}
    doc = Document()
    apply_docx_theme(doc)
    doc.add_heading(name, level=0)

    in_code = False
    code_lang = ""
    code_lines: list[str] = []
    lines = _join_sections(sections).splitlines()
    i = 0

    def flush_code() -> None:
        nonlocal in_code, code_lang, code_lines
        _add_code_block_docx(doc, code_lang, code_lines)
        in_code, code_lang, code_lines = False, "", []

    while i < len(lines):
        raw_line = lines[i]
        stripped = raw_line.strip()
        if stripped.startswith("```"):
            if in_code:
                flush_code()
            else:
                in_code, code_lang, code_lines = True, stripped[3:].strip().lower(), []
            i += 1
            continue
        if in_code:
            code_lines.append(raw_line)
            i += 1
            continue
        if _TABLE_LINE_RE.match(stripped):
            block = [raw_line]
            j = i + 1
            while j < len(lines) and _TABLE_LINE_RE.match(lines[j].strip()):
                block.append(lines[j])
                j += 1
            if len(block) >= 2 and _is_table_sep(_table_cells(block[1])):
                _add_gfm_table_docx(doc, block, assets)
                i = j
                continue
        if not stripped:
            i += 1
            continue
        hm = _HEADING_RE.match(stripped)
        if hm:
            level = min(len(hm.group(1)), 4)
            heading = doc.add_heading("", level=level)
            heading.text = ""
            _add_inline_runs(heading, hm.group(2).strip(), assets)
            i += 1
            continue
        if _HR_RE.match(stripped):
            _add_hr_docx(doc)
            i += 1
            continue
        m = _IMAGE_LINE_RE.match(stripped)
        if m:
            rel = _rel_of(m.group(1))
            data = assets.get(rel)
            if data:
                try:
                    doc.add_picture(io.BytesIO(data), width=Cm(15))
                    i += 1
                    continue
                except Exception:  # noqa: BLE001
                    pass
            para = doc.add_paragraph()
            style_run(para.add_run(f"[画像: {rel}]"), size=Pt(10), color=MUTED, italic=True)
            i += 1
            continue
        qm = _QUOTE_RE.match(stripped)
        if qm:
            para = doc.add_paragraph()
            para.paragraph_format.left_indent = Cm(0.4)
            left_border(para)
            _add_inline_runs(para, qm.group(1), assets, size=Pt(12), color=MUTED)
            for run in para.runs:
                run.italic = True
            i += 1
            continue
        bm = _BULLET_RE.match(raw_line)
        if bm:
            text = bm.group(2)
            task = _TASK_RE.match(text)
            if task:
                mark = "☑" if task.group(1).lower() == "x" else "☐"
                text = f"{mark} {task.group(2)}"
            para = doc.add_paragraph(style="List Bullet")
            para.text = ""
            _add_inline_runs(para, text, assets, size=Pt(12))
            i += 1
            continue
        nm = _NUMBER_RE.match(raw_line)
        if nm:
            para = doc.add_paragraph(style="List Number")
            para.text = ""
            _add_inline_runs(para, nm.group(3), assets, size=Pt(12))
            i += 1
            continue
        para = doc.add_paragraph()
        _add_inline_runs(para, stripped, assets, size=Pt(12))
        i += 1

    if in_code:
        flush_code()
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()

