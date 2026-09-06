"""デジタル庁デザインシステム（DADS）の共通トークンと Word 向け適用。

https://design.digital.go.jp/dads/foundations/color/color-palette/
https://design.digital.go.jp/dads/foundations/typography/
https://www.digital.go.jp/resources/dashboard-guidebook/color-palette/color-code
"""

from __future__ import annotations

from typing import Any

# Solid Gray / Blue（ダッシュボード資料のカラーコードに合わせる）
INK = (0x1A, 0x1A, 0x1A)  # Solid Gray 900
BODY = (0x33, 0x33, 0x33)  # Solid Gray 800
MUTED = (0x76, 0x76, 0x76)  # Solid Gray 536（白地で本文 4.5:1）
RULE = (0xCC, 0xCC, 0xCC)  # Solid Gray 200
PAPER = (0xFF, 0xFF, 0xFF)  # White
SURFACE = (0xF2, 0xF2, 0xF2)  # Solid Gray 50
ACCENT = (0x00, 0x17, 0xC1)  # Blue 900
ON_ACCENT = (0xFF, 0xFF, 0xFF)
FONT = "Noto Sans JP"
FONT_MONO = "Noto Sans Mono"


def hex_of(rgb: tuple[int, int, int]) -> str:
    return f"{rgb[0]:02X}{rgb[1]:02X}{rgb[2]:02X}"


def style_run(
    run: Any,
    *,
    name: str = FONT,
    size: Any = None,
    color: tuple[int, int, int] | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
) -> None:
    """python-docx の Run に和欧とも同じ書体を付ける。"""
    from docx.oxml.ns import qn
    from docx.shared import RGBColor

    run.font.name = name
    if size is not None:
        run.font.size = size
    if color is not None:
        run.font.color.rgb = RGBColor(*color)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        from docx.oxml import OxmlElement

        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    r_fonts.set(qn("w:ascii"), name)
    r_fonts.set(qn("w:hAnsi"), name)
    r_fonts.set(qn("w:eastAsia"), name)
    r_fonts.set(qn("w:cs"), name)


def _set_style_font(
    style: Any,
    *,
    name: str,
    size: Any,
    color: tuple[int, int, int],
    bold: bool | None = None,
) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import RGBColor

    style.font.name = name
    style.font.size = size
    style.font.color.rgb = RGBColor(*color)
    if bold is not None:
        style.font.bold = bold
    r_pr = style.element.get_or_add_rPr()
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    r_fonts.set(qn("w:ascii"), name)
    r_fonts.set(qn("w:hAnsi"), name)
    r_fonts.set(qn("w:eastAsia"), name)
    r_fonts.set(qn("w:cs"), name)


def _bottom_border(p_pr: Any, color: tuple[int, int, int], *, sz: str = "12", space: str = "4") -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    old = p_pr.find(qn("w:pBdr"))
    if old is not None:
        p_pr.remove(old)
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), sz)
    bottom.set(qn("w:space"), space)
    bottom.set(qn("w:color"), hex_of(color))
    border.append(bottom)
    p_pr.append(border)


def _add_page_field(paragraph: Any) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    run = paragraph.add_run()
    style_run(run, name=FONT, size=Pt(9), color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


def apply_docx_theme(doc: Any) -> None:
    """A4・余白・見出し／本文スタイルを DADS に合わせる。"""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt

    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)

    header = section.header.paragraphs[0]
    header.text = ""
    _bottom_border(header._p.get_or_add_pPr(), ACCENT, sz="18", space="1")

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.text = ""
    _add_page_field(footer)

    normal = doc.styles["Normal"]
    _set_style_font(normal, name=FONT, size=Pt(11), color=BODY)
    normal.paragraph_format.line_spacing = 1.7
    normal.paragraph_format.space_after = Pt(10)
    normal.paragraph_format.space_before = Pt(0)

    title = doc.styles["Title"]
    _set_style_font(title, name=FONT, size=Pt(22), color=INK, bold=True)
    title.paragraph_format.space_after = Pt(18)
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.line_spacing = 1.5
    _bottom_border(title.element.get_or_add_pPr(), ACCENT, sz="12", space="6")

    for style_name, size, before in (("Heading 1", 16, 20), ("Heading 2", 14, 16), ("Heading 3", 12, 14)):
        st = doc.styles[style_name]
        _set_style_font(st, name=FONT, size=Pt(size), color=INK, bold=True)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(8)
        st.paragraph_format.line_spacing = 1.5
    _bottom_border(doc.styles["Heading 1"].element.get_or_add_pPr(), RULE, sz="6", space="4")

    try:
        bullet = doc.styles["List Bullet"]
        _set_style_font(bullet, name=FONT, size=Pt(11), color=BODY)
        bullet.paragraph_format.line_spacing = 1.7
        bullet.paragraph_format.space_after = Pt(4)
    except KeyError:
        pass


def shade_paragraph(paragraph: Any, fill: tuple[int, int, int] = SURFACE) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    p_pr = paragraph._p.get_or_add_pPr()
    old = p_pr.find(qn("w:shd"))
    if old is not None:
        p_pr.remove(old)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_of(fill))
    p_pr.append(shd)
