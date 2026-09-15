from shared.navsheet import (
    MARKER,
    NAV_LABEL,
    NavRow,
    NavSheet,
    markdown_to_rows,
    parse_workbook,
    table_to_markdown,
    write_workbook,
)


def test_roundtrip_markdown_and_xlsx():
    rows = [
        NavRow(label="課題は何か", value="予算不足"),
        NavRow(label="パイプ|あり", value="改行\nあり"),
    ]
    md = table_to_markdown(rows)
    parsed = markdown_to_rows(md)
    assert parsed[0].label == "課題は何か"
    assert parsed[0].value == "予算不足"
    assert parsed[1].label == "パイプ|あり"
    assert "改行" in parsed[1].value

    raw = write_workbook(NavSheet(rows=rows, instruction="関係者向けに要約する"))
    sheet = parse_workbook(raw)
    assert sheet.instruction == "関係者向けに要約する"
    assert [r.label for r in sheet.rows] == ["課題は何か", "パイプ|あり"]
    assert sheet.rows[1].value == "改行\nあり"

    import io

    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(raw))
    try:
        ws = wb.active
        assert ws["B1"].value == MARKER
        assert ws["A2"].value == "項目"
        assert ws["B2"].value == "回答"
        assert ws["A3"].value == "課題は何か"
        assert ws["B3"].value == "予算不足"
        assert ws["B4"].value == "改行\nあり"
        assert "<br>" not in (ws["B4"].value or "")
        assert ws["A5"].value == "生成指示"
    finally:
        wb.close()


def test_parse_legacy_markdown_table_cell():
    import io

    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "種別"
    ws["B1"] = MARKER
    ws["A2"] = NAV_LABEL
    ws["B2"] = table_to_markdown(
        [NavRow(label="設問", value="回答<br>続き")],
    )
    ws["A3"] = "生成指示"
    ws["B3"] = "要約する"
    buf = io.BytesIO()
    wb.save(buf)
    sheet = parse_workbook(buf.getvalue())
    assert sheet.instruction == "要約する"
    assert sheet.rows[0].label == "設問"
    assert sheet.rows[0].value == "回答\n続き"
